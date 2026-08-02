from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from workbook_excel_to_json import validate_payload as validate_excel_payload


SCRIPT_DIR = Path(__file__).resolve().parent
JSON_TO_EXCEL = SCRIPT_DIR / "workbook_json_to_excel.py"
EXCEL_TO_JSON = SCRIPT_DIR / "workbook_excel_to_json.py"
VALIDATE_JSON = SCRIPT_DIR / "validate_workbook_json.py"


def sample_payload() -> dict[str, object]:
    return {
        "sheets": [
            {
                "name": "テスト対象",
                "columns": ["テスト対象", "根拠資料", "参照箇所", "参照目的", "備考"],
                "rows": [
                    {
                        "テスト対象": "TicketService.ResolveExpected",
                        "根拠資料": "Spec/APIプロジェクト概要.md",
                        "参照箇所": "2.5",
                        "参照目的": "期待結果の確認",
                        "備考": "",
                    },
                    {
                        "テスト対象": "TicketService.ResolveExpected",
                        "根拠資料": "",
                        "参照箇所": "",
                        "参照目的": "",
                        "備考": "",
                    },
                ],
            },
            {
                "name": "因子と水準",
                "columns": ["因子", "水準", "備考"],
                "rows": [
                    {
                        "因子": "Soup",
                        "水準": "塩",
                        "備考": "",
                    },
                    {
                        "因子": "Soup",
                        "水準": "醤油",
                        "備考": "",
                    },
                    {
                        "因子": "Soup",
                        "水準": "味噌",
                        "備考": "",
                    },
                    {
                        "因子": "NoodleAmount",
                        "水準": "普通",
                        "備考": "Unicodeを保持",
                    },
                    {
                        "因子": "NoodleAmount",
                        "水準": "大盛り",
                        "備考": "Unicodeを保持",
                    },
                ],
            },
            {
                "name": "テストケース",
                "columns": [
                    "ID",
                    "テストの目的",
                    "Soup",
                    "NoodleAmount",
                    "expected",
                    "memo",
                ],
                "rows": [
                    {
                        "ID": "TC-001",
                        "テストの目的": "塩・普通盛りの食券を確認する",
                        "Soup": "塩",
                        "NoodleAmount": "普通",
                        "expected": "食券1",
                        "memo": "仕様: Spec/APIプロジェクト概要.md#2.5",
                    },
                    {
                        "ID": "TC-002",
                        "テストの目的": "味噌・大盛りの未確定結果を確認する",
                        "Soup": "味噌",
                        "NoodleAmount": "大盛り",
                        "expected": "要確認",
                        "memo": "",
                    },
                    {
                        "ID": "TC-003",
                        "テストの目的": "醤油・普通盛りの食券を確認する",
                        "Soup": "醤油",
                        "NoodleAmount": "普通",
                        "expected": "食券3",
                        "memo": "",
                    },
                ],
            },
        ]
    }


class WorkbookJsonToExcelTests(unittest.TestCase):
    def test_round_trip_preserves_workbook_payload(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            testcases = Path(temporary) / "testcases"
            testcases.mkdir()
            source = testcases / "testcase_unicode.json"
            source_payload = sample_payload()
            source_text = json.dumps(source_payload, ensure_ascii=False, indent=2)
            source.write_text(source_text, encoding="utf-8")

            result = subprocess.run(
                [sys.executable, str(JSON_TO_EXCEL), "--input", str(source)],
                check=True,
                capture_output=True,
                text=True,
            )

            workbook_path = source.with_suffix(".xlsx")
            self.assertTrue(workbook_path.is_file())
            self.assertIn("因子と水準=5", result.stdout)
            self.assertEqual(source.read_text(encoding="utf-8"), source_text)

            workbook = load_workbook(workbook_path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["テスト対象", "因子と水準", "テストケース"],
            )
            self.assertEqual(workbook["テスト対象"]["A2"].value, "TicketService.ResolveExpected")
            self.assertIsNone(workbook["テスト対象"]["B3"].value)
            self.assertEqual(workbook["テスト対象"].freeze_panes, "A2")
            self.assertGreater(workbook["テスト対象"].column_dimensions["A"].width, 10)
            self.assertEqual(
                workbook["テスト対象"]["A1"].fill.fgColor.rgb,
                "FFD9EAF7",
            )
            self.assertEqual(workbook["テスト対象"].page_setup.orientation, "landscape")
            self.assertEqual(
                list(
                    workbook["因子と水準"].iter_rows(
                        min_row=1, max_row=6, values_only=True
                    )
                ),
                [
                    ("因子", "水準", "備考"),
                    ("Soup", "塩", None),
                    ("Soup", "醤油", None),
                    ("Soup", "味噌", None),
                    ("NoodleAmount", "普通", "Unicodeを保持"),
                    ("NoodleAmount", "大盛り", "Unicodeを保持"),
                ],
            )
            self.assertEqual(workbook["テストケース"]["F2"].value, "仕様: Spec/APIプロジェクト概要.md#2.5")
            self.assertIsNone(workbook["テストケース"]["F3"].value)

            round_trip = testcases / "testcase_roundtrip.json"
            subprocess.run(
                [
                    sys.executable,
                    str(EXCEL_TO_JSON),
                    "--input",
                    str(workbook_path),
                    "--output",
                    str(round_trip),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertEqual(
                json.loads(round_trip.read_text(encoding="utf-8")), source_payload
            )

    def test_existing_workbook_is_not_overwritten(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            testcases = Path(temporary) / "testcases"
            testcases.mkdir()
            source = testcases / "testcase_existing.json"
            source.write_text(
                json.dumps(sample_payload(), ensure_ascii=False), encoding="utf-8"
            )
            workbook_path = source.with_suffix(".xlsx")
            sentinel = b"existing workbook"
            workbook_path.write_bytes(sentinel)

            result = subprocess.run(
                [sys.executable, str(JSON_TO_EXCEL), "--input", str(source)],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("already exists", result.stderr)
            self.assertEqual(workbook_path.read_bytes(), sentinel)

    def test_invalid_workbook_json_does_not_leave_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            testcases = Path(temporary) / "testcases"
            testcases.mkdir()
            source = testcases / "testcase_invalid.json"
            source.write_text(
                json.dumps({"sheets": []}, ensure_ascii=False), encoding="utf-8"
            )

            result = subprocess.run(
                [sys.executable, str(JSON_TO_EXCEL), "--input", str(source)],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("Workbook sheets must be exactly", result.stderr)
            self.assertFalse(source.with_suffix(".xlsx").exists())

    def test_unresolved_expected_can_be_restored_but_not_implemented(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            testcases = Path(temporary) / "testcases"
            testcases.mkdir()
            source = testcases / "testcase_unresolved.json"
            source.write_text(
                json.dumps(sample_payload(), ensure_ascii=False), encoding="utf-8"
            )

            conversion = subprocess.run(
                [sys.executable, str(JSON_TO_EXCEL), "--input", str(source)],
                check=False,
                capture_output=True,
                text=True,
            )
            validation = subprocess.run(
                [sys.executable, str(VALIDATE_JSON), "--input", str(source)],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(conversion.returncode, 0)
            self.assertTrue(source.with_suffix(".xlsx").is_file())
            self.assertNotEqual(validation.returncode, 0)
            self.assertIn("Unresolved expected values: TC-002", validation.stderr)

    def test_horizontal_factor_levels_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            testcases = Path(temporary) / "testcases"
            testcases.mkdir()
            source = testcases / "testcase_horizontal.json"
            payload = sample_payload()
            payload["sheets"][1] = {
                "name": "因子と水準",
                "columns": ["因子", "水準1", "水準2", "備考"],
                "rows": [
                    {
                        "因子": "Soup",
                        "水準1": "塩",
                        "水準2": "醤油",
                        "備考": "",
                    }
                ],
            }
            source.write_text(
                json.dumps(payload, ensure_ascii=False), encoding="utf-8"
            )

            result = subprocess.run(
                [sys.executable, str(JSON_TO_EXCEL), "--input", str(source)],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("columns must be exactly", result.stderr)
            self.assertFalse(source.with_suffix(".xlsx").exists())

    def test_duplicated_factor_and_level_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            testcases = Path(temporary) / "testcases"
            testcases.mkdir()
            source = testcases / "testcase_duplicate_level.json"
            payload = sample_payload()
            payload["sheets"][1]["rows"].append(
                {"因子": "Soup", "水準": "塩", "備考": "重複"}
            )
            source.write_text(
                json.dumps(payload, ensure_ascii=False), encoding="utf-8"
            )

            result = subprocess.run(
                [sys.executable, str(VALIDATE_JSON), "--input", str(source)],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("Duplicated 因子 and 水準: Soup / 塩", result.stderr)

    def test_excel_conversion_rejects_duplicated_factor_and_level(self) -> None:
        payload = sample_payload()
        payload["sheets"][1]["rows"].append(
            {"因子": "Soup", "水準": "塩", "備考": "重複"}
        )

        with self.assertRaisesRegex(
            ValueError,
            "Duplicated 因子 and 水準: Soup / 塩",
        ):
            validate_excel_payload(payload)

    def test_empty_test_target_is_rejected(self) -> None:
        payload = sample_payload()
        payload["sheets"][0]["rows"][0]["テスト対象"] = ""

        with self.assertRaisesRegex(ValueError, "empty テスト対象"):
            validate_excel_payload(payload)

    def test_empty_test_purpose_is_rejected(self) -> None:
        payload = sample_payload()
        payload["sheets"][2]["rows"][0]["テストの目的"] = ""

        with self.assertRaisesRegex(ValueError, "empty テストの目的"):
            validate_excel_payload(payload)

    def test_additional_sheet_is_rejected(self) -> None:
        payload = sample_payload()
        payload["sheets"].append({"name": "補助", "columns": [], "rows": []})

        with self.assertRaisesRegex(ValueError, "Workbook sheets must be exactly"):
            validate_excel_payload(payload)

    def test_undefined_level_is_rejected(self) -> None:
        payload = sample_payload()
        payload["sheets"][2]["rows"][0]["Soup"] = "豚骨"

        with self.assertRaisesRegex(ValueError, "uses undefined level: Soup / 豚骨"):
            validate_excel_payload(payload)

    def test_unused_level_is_rejected(self) -> None:
        payload = sample_payload()
        payload["sheets"][2]["rows"] = payload["sheets"][2]["rows"][:2]

        with self.assertRaisesRegex(ValueError, "Unused factor levels: Soup / 醤油"):
            validate_excel_payload(payload)


if __name__ == "__main__":
    unittest.main()
