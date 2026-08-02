from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from validate_workbook_json import validate_payload


def sheet_to_payload(ws: Any) -> dict[str, Any]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"name": ws.title, "columns": [], "rows": []}
    slots: list[tuple[int, str]] = []
    columns: list[str] = []
    for index, value in enumerate(rows[0]):
        name = "" if value is None else str(value).strip()
        if not name:
            continue
        if name in columns:
            raise ValueError(f"Sheet {ws.title!r} has duplicated column: {name}")
        slots.append((index, name))
        columns.append(name)
    data_rows: list[dict[str, str]] = []
    for values in rows[1:]:
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        item = {
            name: "" if index >= len(values) or values[index] is None else str(values[index])
            for index, name in slots
        }
        if any(value.strip() for value in item.values()):
            data_rows.append(item)
    return {"name": ws.title, "columns": columns, "rows": data_rows}


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert a test matrix workbook to workbook JSON.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output")
    args = parser.parse_args()
    input_path = Path(args.input)
    if not input_path.is_file():
        raise ValueError(f"Input workbook not found: {input_path}")
    if input_path.parent.name != "testcases" or not input_path.name.startswith("testcase_"):
        raise ValueError("Input must be named testcases/testcase_*.xlsx")
    output_path = Path(args.output) if args.output else input_path.with_suffix(".json")
    if output_path.parent.name != "testcases" or not output_path.name.startswith("testcase_") or output_path.suffix.lower() != ".json":
        raise ValueError("Output must be named testcases/testcase_*.json")
    workbook = load_workbook(input_path, data_only=False)
    try:
        payload = {
            "sheets": [
                sheet_to_payload(workbook[name]) for name in workbook.sheetnames
            ]
        }
    finally:
        workbook.close()
    validate_payload(payload)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    counts = ", ".join(f"{sheet['name']}={len(sheet['rows'])}" for sheet in payload["sheets"])
    print(f"Converted {len(payload['sheets'])} sheets ({counts}) -> {output_path}")


if __name__ == "__main__":
    main()
