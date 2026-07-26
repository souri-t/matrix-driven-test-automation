from __future__ import annotations

import argparse
import json
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

from validate_workbook_json import validate


def validate_paths(input_path: Path, output_path: Path) -> None:
    if not input_path.is_file():
        raise ValueError(f"Input JSON not found: {input_path}")
    if (
        input_path.parent.name != "testcases"
        or not input_path.name.startswith("testcase_")
        or input_path.suffix.lower() != ".json"
    ):
        raise ValueError("Input must be named testcases/testcase_*.json")
    if output_path != input_path.with_suffix(".xlsx"):
        raise ValueError("Output must have the same path and stem as the input JSON")
    if output_path.exists():
        raise FileExistsError(f"Output workbook already exists: {output_path}")


def write_workbook(payload: dict[str, Any], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in payload["sheets"]:
        worksheet = workbook.create_sheet(title=sheet["name"])
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        columns = sheet["columns"]
        worksheet.append(columns)
        for row in sheet["rows"]:
            worksheet.append([row.get(column, "") for column in columns])
        for cell in worksheet[1]:
            cell.fill = PatternFill(fill_type="solid", fgColor="FFD9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(columns))}{worksheet.max_row}"
            )
        for index, column in enumerate(columns, start=1):
            values = [column, *(row.get(column, "") for row in sheet["rows"])]
            width = min(
                40,
                max(10, max(_display_width(str(value)) for value in values) + 2),
            )
            worksheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(output_path)


def _display_width(value: str) -> int:
    return sum(
        2 if unicodedata.east_asian_width(character) in {"F", "W", "A"} else 1
        for character in value
    )


def verify_workbook(payload: dict[str, Any], output_path: Path) -> None:
    workbook = load_workbook(output_path, data_only=False, read_only=True)
    try:
        expected_names = [sheet["name"] for sheet in payload["sheets"]]
        if workbook.sheetnames != expected_names:
            raise ValueError(
                f"Sheet order mismatch: expected {expected_names}, got {workbook.sheetnames}"
            )
        for sheet in payload["sheets"]:
            worksheet = workbook[sheet["name"]]
            values = list(worksheet.iter_rows(values_only=True))
            actual_columns = [
                "" if value is None else str(value)
                for value in (values[0] if values else ())
            ]
            if actual_columns != sheet["columns"]:
                raise ValueError(f"Column mismatch in sheet {sheet['name']!r}")
            actual_rows = [
                {
                    column: ""
                    if index >= len(row) or row[index] is None
                    else str(row[index])
                    for index, column in enumerate(sheet["columns"])
                }
                for row in values[1:]
            ]
            expected_rows = [
                {column: str(row.get(column, "")) for column in sheet["columns"]}
                for row in sheet["rows"]
            ]
            if actual_rows != expected_rows:
                raise ValueError(f"Row mismatch in sheet {sheet['name']!r}")
    finally:
        workbook.close()


def convert(input_path: Path) -> tuple[Path, list[tuple[str, int]]]:
    output_path = input_path.with_suffix(".xlsx")
    validate_paths(input_path, output_path)
    validate(input_path)
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    try:
        output_path.touch(exist_ok=False)
    except FileExistsError as error:
        raise FileExistsError(
            f"Output workbook already exists: {output_path}"
        ) from error
    try:
        write_workbook(payload, output_path)
        verify_workbook(payload, output_path)
    except Exception:
        if output_path.exists():
            output_path.unlink()
        raise
    counts = [(sheet["name"], len(sheet["rows"])) for sheet in payload["sheets"]]
    return output_path, counts


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert workbook JSON to a same-named Excel workbook."
    )
    parser.add_argument("--input", required=True)
    args = parser.parse_args()
    output_path, counts = convert(Path(args.input))
    summary = ", ".join(f"{name}={count}" for name, count in counts)
    print(f"Converted {len(counts)} sheets ({summary}) -> {output_path}")


if __name__ == "__main__":
    main()
