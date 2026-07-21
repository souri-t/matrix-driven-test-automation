from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_SHEETS = ("因子と水準", "テストケース")
RESERVED_TEST_COLUMNS = {"ID", "expected", "memo"}


def normalize_header(value: Any) -> str:
    return "" if value is None else str(value).strip()


def is_empty_row(values: list[Any]) -> bool:
    return all(value is None or str(value).strip() == "" for value in values)


def sheet_to_payload(ws: Any) -> dict[str, Any]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"name": ws.title, "columns": [], "rows": []}

    header_slots: list[tuple[int, str]] = []
    columns: list[str] = []
    for index, raw_header in enumerate(rows[0]):
        header = normalize_header(raw_header)
        if not header:
            continue
        if header in columns:
            raise ValueError(f"Sheet {ws.title!r} has duplicated column: {header}")
        header_slots.append((index, header))
        columns.append(header)

    data_rows: list[dict[str, Any]] = []
    for values in rows[1:]:
        row_values = list(values)
        if is_empty_row(row_values):
            continue
        item = {
            name: "" if index >= len(row_values) or row_values[index] is None else str(row_values[index])
            for index, name in header_slots
        }
        if any(value.strip() for value in item.values()):
            data_rows.append(item)

    return {"name": ws.title, "columns": columns, "rows": data_rows}


def validate_payload(payload: dict[str, Any]) -> None:
    sheets = {sheet["name"]: sheet for sheet in payload["sheets"]}
    missing_sheets = [name for name in REQUIRED_SHEETS if name not in sheets]
    if missing_sheets:
        raise ValueError(f"Missing required sheets: {', '.join(missing_sheets)}")

    factor_columns = sheets["因子と水準"]["columns"]
    missing_factor_columns = [name for name in ("因子", "水準1", "備考") if name not in factor_columns]
    if missing_factor_columns:
        raise ValueError(f"Sheet '因子と水準' is missing columns: {', '.join(missing_factor_columns)}")

    testcase_columns = sheets["テストケース"]["columns"]
    missing_columns = [name for name in ("ID", "expected", "memo") if name not in testcase_columns]
    if missing_columns:
        raise ValueError(f"Sheet 'テストケース' is missing columns: {', '.join(missing_columns)}")
    if not [name for name in testcase_columns if name not in RESERVED_TEST_COLUMNS]:
        raise ValueError("Sheet 'テストケース' must contain at least one factor column")

    seen_ids: set[str] = set()
    for index, row in enumerate(sheets["テストケース"]["rows"], start=2):
        case_id = row["ID"].strip()
        if not case_id:
            raise ValueError(f"Sheet 'テストケース' row {index} has an empty ID")
        if case_id in seen_ids:
            raise ValueError(f"Sheet 'テストケース' has duplicated ID: {case_id}")
        seen_ids.add(case_id)
        if not row["expected"].strip():
            raise ValueError(f"Sheet 'テストケース' row {index} has an empty expected")


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert a test matrix workbook to workbook JSON.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.is_file():
        raise ValueError(f"Input workbook not found: {input_path}")
    if input_path.parent.name != "testcases" or not input_path.name.startswith("testcase_"):
        raise ValueError("Input must be named testcases/testcase_*.xlsx")
    output_path = Path(args.output) if args.output else input_path.with_suffix(".json")
    if (
        output_path.parent.name != "testcases"
        or not output_path.name.startswith("testcase_")
        or output_path.suffix.lower() != ".json"
    ):
        raise ValueError("Output must be named testcases/testcase_*.json")

    workbook = load_workbook(input_path, data_only=False)
    payload = {"sheets": [sheet_to_payload(workbook[name]) for name in workbook.sheetnames]}
    validate_payload(payload)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    counts = ", ".join(f"{sheet['name']}={len(sheet['rows'])}" for sheet in payload["sheets"])
    print(f"Converted {len(payload['sheets'])} sheets ({counts}) -> {output_path}")


if __name__ == "__main__":
    main()
